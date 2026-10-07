"""Tests for the CSV/XLSX loader: delimiters, encodings, typing, caps."""

from __future__ import annotations

import io
from datetime import datetime

import pytest
from openpyxl import Workbook

from src.services import table_loader
from src.services.table_loader import load_table, parse_date, parse_number
from src.services.tables import TableRefused


def test_a_comma_csv_is_typed_per_column():
    csv = b"region,units,sold\nEU,10,2026-01-01\nUS,\"1,200\",2026-01-02\nAPAC,7,2026-02-01\n"

    parsed = load_table("sales.csv", csv)

    assert parsed.columns == ["region", "units", "sold"]
    assert parsed.kinds == ["text", "number", "date"]
    assert parsed.data[0] == ["EU", "US", "APAC"]
    assert parsed.data[1] == [10.0, 1200.0, 7.0]
    assert parsed.data[2][0] == datetime(2026, 1, 1)
    assert parsed.row_count == 3 and parsed.sheet is None


def test_semicolon_and_tab_delimiters_are_sniffed():
    assert load_table("a.csv", b"a;b\n1;2\n").columns == ["a", "b"]
    assert load_table("a.csv", b"a\tb\n1\t2\n").columns == ["a", "b"]


def test_bom_and_cp1252_are_decoded():
    assert load_table("a.csv", b"\xef\xbb\xbfname\nx\n").columns == ["name"]
    assert load_table("a.csv", "city\nZ\xfcrich\n".encode("cp1252")).data[0] == ["Z\xfcrich"]


def test_blank_and_duplicate_headers_are_renamed():
    parsed = load_table("a.csv", b"x,,x,x\n1,2,3,4\n")

    assert parsed.columns == ["x", "column_2", "x_2", "x_3"]


def test_blank_rows_are_skipped_and_short_rows_padded_and_reported():
    parsed = load_table("a.csv", b"a,b\n1,2\n\n3\n")

    assert parsed.row_count == 2
    assert parsed.data[1] == [2.0, None]
    assert any("1 rows" in note and "different number of cells" in note for note in parsed.notes)


def test_numbers_accept_thousands_separators_and_percent():
    parsed = load_table("a.csv", b'v\n"1,234.5"\n12%\n-3\n')

    assert parsed.kinds == ["number"]
    assert parsed.data[0] == [1234.5, 12.0, -3.0]


def test_a_mostly_numeric_column_keeps_the_numbers_and_reports_the_rest():
    rows = "\n".join(["1"] * 19 + ["oops"])
    parsed = load_table("a.csv", f"v\n{rows}\n".encode())

    assert parsed.kinds == ["number"]
    assert parsed.data[0].count(None) == 1
    assert any("1 cells" in note and "'v'" in note for note in parsed.notes)


def test_a_half_numeric_column_stays_text():
    parsed = load_table("a.csv", b"v\n1\nabc\n2\nxyz\n")

    assert parsed.kinds == ["text"]
    assert parsed.data[0] == ["1", "abc", "2", "xyz"]


def test_an_empty_column_is_text_with_nothing_in_it():
    parsed = load_table("a.csv", b"a,b\n1,\n2,\n")

    assert parsed.kinds == ["number", "text"]
    assert parsed.data[1] == [None, None]


def test_unsupported_empty_and_header_only_files_are_refused():
    with pytest.raises(TableRefused, match=r"Only \.csv and \.xlsx"):
        load_table("a.txt", b"x")
    with pytest.raises(TableRefused, match="header row and at least one data row"):
        load_table("a.csv", b"a,b\n")
    with pytest.raises(TableRefused, match="header row and at least one data row"):
        load_table("a.csv", b"")


def test_the_row_and_column_caps_refuse_the_file(monkeypatch):
    monkeypatch.setattr(table_loader, "MAX_ROWS", 3)
    with pytest.raises(TableRefused, match="more than 3 rows"):
        load_table("a.csv", b"a\n1\n2\n3\n4\n")
    monkeypatch.setattr(table_loader, "MAX_COLUMNS", 2)
    with pytest.raises(TableRefused, match="more than 2 columns"):
        load_table("a.csv", b"a,b,c\n1,2,3\n")


def test_a_file_over_the_byte_cap_is_refused(monkeypatch):
    monkeypatch.setattr(table_loader, "MAX_FILE_BYTES", 10)
    with pytest.raises(TableRefused, match="larger than"):
        load_table("a.csv", b"a,b\n1,2\n3,4\n")


def workbook_bytes(sheets: dict[str, list[list]]) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_xlsx_uses_the_first_sheet_with_data_and_keeps_its_name():
    content = workbook_bytes(
        {
            "Empty": [],
            "Sales": [["region", "units", "sold"], ["EU", 10, datetime(2026, 1, 1)], ["US", 20, datetime(2026, 1, 2)]],
            "Other": [["x"], [1]],
        }
    )

    parsed = load_table("book.xlsx", content)

    assert parsed.sheet == "Sales"
    assert parsed.columns == ["region", "units", "sold"]
    assert parsed.kinds == ["text", "number", "date"]
    assert parsed.data[1] == [10.0, 20.0]


def test_a_corrupt_xlsx_is_refused():
    with pytest.raises(TableRefused, match="Could not read this Excel file"):
        load_table("book.xlsx", b"not a zip")


def test_a_workbook_with_no_data_is_refused():
    with pytest.raises(TableRefused, match="no data"):
        load_table("book.xlsx", workbook_bytes({"Empty": []}))


def test_a_zip_bomb_workbook_is_refused(monkeypatch):
    monkeypatch.setattr(table_loader, "MAX_UNZIPPED_BYTES", 10)
    with pytest.raises(TableRefused, match="too large to read"):
        load_table("book.xlsx", workbook_bytes({"S": [["a"], [1]]}))


def test_parse_number_and_parse_date():
    assert parse_number("1,234") == 1234.0 and parse_number(" 7 ") == 7.0 and parse_number(True) is None
    assert parse_number("1,23") is None and parse_number("abc") is None and parse_number(3) == 3.0
    assert parse_date("2026-01-31") == datetime(2026, 1, 31)
    assert parse_date("2026-01-31 10:30:00") == datetime(2026, 1, 31, 10, 30)
    assert parse_date("31/01/2026") is None and parse_date(None) is None


def test_blank_rows_do_not_count_toward_the_row_cap(monkeypatch):
    monkeypatch.setattr(table_loader, "MAX_ROWS", 3)

    parsed = load_table("a.csv", b"a\n1\n2\n" + b"\n" * 20)
    assert parsed.row_count == 2

    content = workbook_bytes({"S": [["a"], [1], [2]] + [[None]] * 20})
    assert load_table("a.xlsx", content).row_count == 2


def test_an_absurdly_long_digit_string_is_not_a_number():
    assert parse_number("9" * 400) is None
    assert parse_number(float("inf")) is None


def test_a_duplicate_header_rename_never_collides_with_an_existing_name():
    parsed = load_table("a.csv", b"x,x,x_2\n1,2,3\n")

    assert len(set(parsed.columns)) == 3
    assert parsed.columns[:2] == ["x", "x_2"]
