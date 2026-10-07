"""Turns an uploaded CSV or XLSX file into a typed `ParsedTable`.

Stdlib `csv` plus `openpyxl` in read-only mode - no pandas. Everything is
refused with a `TableRefused` whose message is safe to show to the caller.

Typing is per column, after the whole file is read: number when at least 95%
of the non-empty cells parse as numbers, else date when at least 95% parse as
ISO dates, else text. A cell that fails the chosen type becomes empty and is
counted in the table's `notes`, so a stray "n/a" never silently skews a total
without the caller being told.

Callers run `load_table` off the event loop (it is blocking, CPU-bound work).
"""

from __future__ import annotations

import csv
import io
import re
import zipfile
from datetime import date, datetime
from typing import Any

import openpyxl

from src.services.tables import MAX_COLUMNS, MAX_FILE_BYTES, MAX_ROWS, ParsedTable, TableRefused

MAX_ZIP_ENTRIES = 5_000
MAX_UNZIPPED_BYTES = 100 * 1024 * 1024
TYPE_THRESHOLD = 0.95
_DELIMITERS = (",", ";", "\t")
_NUMBER = re.compile(r"[+-]?(?:\d+|\d{1,3}(?:,\d{3})+)(?:\.\d+)?%?|[+-]?\.\d+%?")


def load_table(filename: str, content: bytes) -> ParsedTable:
    """Parses `content` (a .csv or .xlsx file named `filename`) into typed columns."""
    if len(content) > MAX_FILE_BYTES:
        raise TableRefused(f"The file is larger than {MAX_FILE_BYTES // (1024 * 1024)} MB.")
    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if suffix == "csv":
        sheet, rows = None, _csv_rows(content)
    elif suffix == "xlsx":
        sheet, rows = _xlsx_rows(content)
    else:
        raise TableRefused("Only .csv and .xlsx files can be analysed.")
    return _build(rows, sheet)


def parse_number(value: Any) -> float | None:
    """A number from a cell, or None. "1,234.5" and "12%" parse; a bool never does."""
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if not isinstance(value, str):
        return None
    text = value.strip()
    if not _NUMBER.fullmatch(text):
        return None
    return float(text.rstrip("%").replace(",", ""))


def parse_date(value: Any) -> datetime | None:
    """A datetime from a cell (a date, a datetime or an ISO string), or None."""
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    if not isinstance(value, str):
        return None
    try:
        return datetime.fromisoformat(value.strip())
    except ValueError:
        return None


def _filled(cell: Any) -> bool:
    return cell is not None and not (isinstance(cell, str) and cell.strip() == "")


def _csv_rows(content: bytes) -> list[list[Any]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("cp1252", errors="replace")
    first_line = text.split("\n", 1)[0]
    delimiter = max(_DELIMITERS, key=first_line.count)
    rows: list[list[Any]] = []
    try:
        for row in csv.reader(io.StringIO(text, newline=""), delimiter=delimiter):
            rows.append(row)
            if len(rows) > MAX_ROWS + 1:
                raise TableRefused(f"The file has more than {MAX_ROWS:,} rows.")
    except csv.Error as error:
        raise TableRefused("Could not read this CSV file.") from error
    return rows


def _xlsx_rows(content: bytes) -> tuple[str, list[list[Any]]]:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ZIP_ENTRIES or sum(entry.file_size for entry in entries) > MAX_UNZIPPED_BYTES:
                raise TableRefused("This workbook is too large to read.")
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except TableRefused:
        raise
    except Exception as error:  # noqa: BLE001 - BadZipFile, KeyError and openpyxl's own errors all mean "corrupt file"
        raise TableRefused("Could not read this Excel file.") from error
    try:
        for sheet in workbook.worksheets:
            rows: list[list[Any]] = []
            for row in sheet.iter_rows(values_only=True):
                rows.append(list(row))
                if len(rows) > MAX_ROWS + 1:
                    raise TableRefused(f"The sheet has more than {MAX_ROWS:,} rows.")
            if any(_filled(cell) for row in rows for cell in row):
                return sheet.title, rows
    finally:
        workbook.close()
    raise TableRefused("The workbook has no data.")


def _build(rows: list[list[Any]], sheet: str | None) -> ParsedTable:
    rows = [row for row in rows if any(_filled(cell) for cell in row)]
    if len(rows) < 2:
        raise TableRefused("The file needs a header row and at least one data row.")
    header, body = rows[0], rows[1:]
    width = len(header)
    while width and not _filled(header[width - 1]):
        width -= 1
    if width > MAX_COLUMNS:
        raise TableRefused(f"The file has more than {MAX_COLUMNS} columns.")
    columns = _names(header[:width])
    raw: list[list[Any]] = [[] for _ in range(width)]
    ragged = 0
    for row in body:
        if len(row) < width or any(_filled(cell) for cell in row[width:]):
            ragged += 1
        for index in range(width):
            raw[index].append(row[index] if index < len(row) else None)
    notes: list[str] = []
    if ragged:
        notes.append(f"{ragged} rows had a different number of cells than the header; they were padded or cut.")
    kinds: list[str] = []
    data: list[list[Any]] = []
    for name, values in zip(columns, raw, strict=True):
        kind, typed, bad = _type_column(values)
        kinds.append(kind)
        data.append(typed)
        if bad:
            notes.append(f"Column {name!r}: {bad} cells did not look like {kind} values and were left empty.")
    return ParsedTable(columns=columns, kinds=kinds, data=data, row_count=len(body), sheet=sheet, notes=notes)


def _names(cells: list[Any]) -> list[str]:
    seen: dict[str, int] = {}
    names: list[str] = []
    for position, cell in enumerate(cells, start=1):
        base = str(cell).strip() if _filled(cell) else f"column_{position}"
        seen[base] = seen.get(base, 0) + 1
        names.append(base if seen[base] == 1 else f"{base}_{seen[base]}")
    return names


def _type_column(values: list[Any]) -> tuple[str, list[Any], int]:
    filled = sum(1 for value in values if _filled(value))
    if not filled:
        return "text", [None] * len(values), 0
    for kind, convert in (("number", parse_number), ("date", parse_date)):
        converted = [convert(value) if _filled(value) else None for value in values]
        good = sum(1 for value in converted if value is not None)
        if good >= TYPE_THRESHOLD * filled:
            return kind, converted, filled - good
    return "text", [str(value).strip() if _filled(value) else None for value in values], 0
