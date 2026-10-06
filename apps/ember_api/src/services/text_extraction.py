"""Plain text from a file attached to a chat question, so it can go into
the question as context (port of chat_app/src/services/text_extraction.py).
The agents are text-only, so the file itself goes no further than this.

Parsing is CPU work in third-party libraries; extract_text_async runs it in
a worker thread so the event loop keeps serving other requests.
"""

from __future__ import annotations

import asyncio
import io
import zipfile
from dataclasses import dataclass

MAX_FILE_BYTES = 15 * 1024 * 1024  # a large document, not a video dropped by mistake
MAX_TEXT_CHARS = 20_000  # about 5k tokens: one attachment mustn't take over the context

# A small file can hide a lot of work: a zip that expands to gigabytes, or a
# sheet of millions of empty rows. Every reader stops at these.
MAX_UNZIPPED_BYTES = 100 * 1024 * 1024
MAX_ZIP_ENTRIES = 5_000
MAX_PDF_PAGES = 500
MAX_SHEET_ROWS = 200_000

PLAIN_TEXT_SUFFIXES = frozenset({
    ".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".log", ".ini", ".cfg", ".conf",
    ".yaml", ".yml", ".xml", ".html", ".htm", ".css", ".sql",
    ".py", ".js", ".ts", ".tsx", ".jsx", ".vue", ".java", ".c", ".h", ".cpp", ".hpp", ".cs", ".go",
    ".rb", ".php", ".sh", ".bat", ".ps1", ".rs", ".kt", ".swift", ".toml",
})  # fmt: skip


class ExtractionError(Exception):
    """Safe to show the user as is: unsupported type, empty, too big, or
    unreadable."""


@dataclass(frozen=True)
class ExtractedText:
    filename: str
    text: str
    char_count: int
    truncated: bool


async def extract_text_async(filename: str, content: bytes) -> ExtractedText:
    return await asyncio.to_thread(extract_text, filename, content)


def extract_text(filename: str, content: bytes) -> ExtractedText:
    if not content:
        raise ExtractionError("The file is empty.")
    if len(content) > MAX_FILE_BYTES:
        raise ExtractionError(f"The file is too large - the limit is {MAX_FILE_BYTES // (1024 * 1024)} MB.")

    suffix = _suffix(filename)
    if suffix == ".pdf":
        text = _pdf(content)
    elif suffix == ".docx":
        text = _docx(content)
    elif suffix in (".xlsx", ".xlsm"):
        text = _xlsx(content)
    elif suffix in PLAIN_TEXT_SUFFIXES:
        text = _plain(content)
    else:
        raise ExtractionError(
            f"Can't read '{suffix or 'extensionless'}' files - try a text, code, .pdf, .docx or .xlsx file."
        )

    text = text.strip()
    if not text:
        raise ExtractionError("No readable text found in this file.")
    truncated = len(text) > MAX_TEXT_CHARS
    if truncated:
        text = text[:MAX_TEXT_CHARS]
    return ExtractedText(filename=filename, text=text, char_count=len(text), truncated=truncated)


def _enough(parts: list[str]) -> bool:
    """True once the text read so far already fills the attachment, so the rest of the file need not be read."""
    return sum(len(p) for p in parts) >= MAX_TEXT_CHARS


def _refuse_zip_bomb(content: bytes) -> None:
    """Office files are zip archives: refuse one that claims to expand beyond
    any real document, before a library unpacks it."""
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
    except zipfile.BadZipFile:
        return  # the library reports an unreadable file in its own words
    if len(entries) > MAX_ZIP_ENTRIES or sum(e.file_size for e in entries) > MAX_UNZIPPED_BYTES:
        raise ExtractionError("This file is far larger than it looks once unpacked, so it was not opened.")


def _suffix(filename: str) -> str:
    dot = filename.rfind(".")
    return filename[dot:].lower() if dot != -1 else ""


def _plain(content: bytes) -> str:
    text = content.decode("utf-8", errors="replace")
    # A binary file "decodes" too, into replacement characters.
    if text.count("�") > len(text) * 0.05:
        raise ExtractionError("This doesn't look like a text file - it isn't valid UTF-8.")
    return text


def _pdf(content: bytes) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(content))
    except Exception as error:  # noqa: BLE001 - any parse failure is one user-facing message
        raise ExtractionError(f"Could not read this PDF: {error}") from error
    pages: list[str] = []
    for number, page in enumerate(reader.pages):
        if number >= MAX_PDF_PAGES or _enough(pages):
            break
        try:
            pages.append(page.extract_text() or "")
        except Exception:  # noqa: BLE001, S112 - one broken page shouldn't lose the rest
            continue
    return "\n\n".join(pages)


def _xlsx(content: bytes) -> str:
    from openpyxl import load_workbook

    _refuse_zip_bomb(content)
    try:
        # data_only: a formula's last computed value, as the sheet shows it.
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as error:  # noqa: BLE001
        raise ExtractionError(f"Could not read this Excel file: {error}") from error
    several = len(workbook.sheetnames) > 1
    parts: list[str] = []
    for sheet in workbook.worksheets:
        if _enough(parts):
            break
        rows: list[str] = []
        for scanned, row in enumerate(sheet.iter_rows(values_only=True)):
            if scanned >= MAX_SHEET_ROWS or _enough([*parts, *rows]):
                break
            # Blank cells stay as empty columns so later values don't shift left.
            cells = ["" if cell is None else str(cell) for cell in row]
            while cells and not cells[-1]:
                cells.pop()
            if cells:
                rows.append(" | ".join(cells))
        if rows:
            parts.append((f"# {sheet.title}\n" if several else "") + "\n".join(rows))
    workbook.close()
    return "\n\n".join(parts)


def _docx(content: bytes) -> str:
    from docx import Document

    _refuse_zip_bomb(content)
    try:
        document = Document(io.BytesIO(content))
    except Exception as error:  # noqa: BLE001
        raise ExtractionError(f"Could not read this Word document: {error}") from error
    parts: list[str] = []
    for paragraph in document.paragraphs:
        if _enough(parts):
            break
        if paragraph.text:
            parts.append(paragraph.text)
    for table in document.tables:
        for row in table.rows:
            if _enough(parts):
                return "\n".join(parts)
            cells = [cell.text for cell in row.cells if cell.text]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)
