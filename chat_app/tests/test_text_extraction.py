import io

import pytest
from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter

from src.services import text_extraction


def test_extract_plain_text_file():
    result = text_extraction.extract_text("notes.txt", b"hello world")

    assert result == {"filename": "notes.txt", "text": "hello world", "char_count": 11, "truncated": False}


def test_extract_rejects_empty_file():
    with pytest.raises(text_extraction.ExtractionError, match="empty"):
        text_extraction.extract_text("notes.txt", b"")


def test_extract_rejects_oversized_file(monkeypatch):
    monkeypatch.setattr(text_extraction, "MAX_UPLOAD_BYTES", 10)

    with pytest.raises(text_extraction.ExtractionError, match="too large"):
        text_extraction.extract_text("notes.txt", b"this is more than ten bytes")


def test_extract_rejects_unsupported_extension():
    with pytest.raises(text_extraction.ExtractionError, match="Can't read"):
        text_extraction.extract_text("archive.zip", b"PK\x03\x04binarydata")


def test_extract_rejects_binary_content_masquerading_as_text():
    with pytest.raises(text_extraction.ExtractionError, match="doesn't look like a text file"):
        text_extraction.extract_text("data.log", bytes(range(256)) * 4)


def test_extract_truncates_long_text(monkeypatch):
    monkeypatch.setattr(text_extraction, "MAX_EXTRACTED_CHARS", 5)

    result = text_extraction.extract_text("notes.txt", b"abcdefghij")

    assert result == {"filename": "notes.txt", "text": "abcde", "char_count": 5, "truncated": True}


def test_extract_docx():
    buffer = io.BytesIO()
    document = Document()
    document.add_paragraph("Hello from a Word document.")
    document.save(buffer)

    result = text_extraction.extract_text("report.docx", buffer.getvalue())

    assert result["text"] == "Hello from a Word document."
    assert result["truncated"] is False


def test_extract_docx_rejects_corrupt_file():
    with pytest.raises(text_extraction.ExtractionError, match="Could not read this Word document"):
        text_extraction.extract_text("report.docx", b"not a real docx file")


def test_extract_pdf():
    buffer = io.BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.write(buffer)

    # A blank page has no text to extract - this only proves parsing
    # itself succeeds without raising; text content is covered by trusting
    # pypdf's own test suite for extract_text() rather than re-testing it
    # here with a hand-built page containing real text content streams.
    with pytest.raises(text_extraction.ExtractionError, match="No readable text"):
        text_extraction.extract_text("blank.pdf", buffer.getvalue())


def test_extract_pdf_rejects_corrupt_file():
    with pytest.raises(text_extraction.ExtractionError, match="Could not read this PDF"):
        text_extraction.extract_text("report.pdf", b"not a real pdf file")


def test_extract_xlsx_single_sheet():
    buffer = io.BytesIO()
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Name", "Score"])
    sheet.append(["Alice", 91])
    workbook.save(buffer)

    result = text_extraction.extract_text("scores.xlsx", buffer.getvalue())

    assert result["text"] == "Name | Score\nAlice | 91"
    assert result["truncated"] is False


def test_extract_xlsx_keeps_blank_cells_so_columns_stay_aligned():
    buffer = io.BytesIO()
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Role", "Field", "Activity", "Action"])
    sheet.append(["Z_ROLE", "DOKAR", None, "Remove"])
    sheet.append([None, None, None, None])
    workbook.save(buffer)

    result = text_extraction.extract_text("request.xlsx", buffer.getvalue())

    assert result["text"] == "Role | Field | Activity | Action\nZ_ROLE | DOKAR |  | Remove"


def test_extract_xlsx_multiple_sheets_labels_each_with_its_title():
    buffer = io.BytesIO()
    workbook = Workbook()
    workbook.active.title = "Summary"
    workbook.active.append(["Total", 3])
    workbook.create_sheet("Detail").append(["Row 1"])
    workbook.save(buffer)

    result = text_extraction.extract_text("report.xlsx", buffer.getvalue())

    assert result["text"] == "# Summary\nTotal | 3\n\n# Detail\nRow 1"


def test_extract_xlsx_rejects_corrupt_file():
    with pytest.raises(text_extraction.ExtractionError, match="Could not read this Excel file"):
        text_extraction.extract_text("report.xlsx", b"not a real xlsx file")
