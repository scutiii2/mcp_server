from __future__ import annotations

import base64
import io

from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook

from src.services.text_extraction import MAX_TEXT_CHARS
from tests.test_registration import as_admin


def attach(client: TestClient, filename: str, content: bytes):
    return client.post(
        "/api/attachments/text", json={"filename": filename, "data": base64.b64encode(content).decode()}
    )


def test_text_and_code_files(client: TestClient) -> None:
    assert attach(client, "notes.txt", b"hi").status_code == 401
    as_admin(client)

    response = attach(client, "C:\\Users\\me\\notes.md", "# Title\n\nSome text ✓\n".encode())
    assert response.status_code == 200, response.text
    assert response.json() == {"filename": "notes.md", "text": "# Title\n\nSome text ✓", "char_count": 20, "truncated": False}

    long = attach(client, "big.log", b"x" * (MAX_TEXT_CHARS + 50)).json()
    assert (long["char_count"], long["truncated"]) == (MAX_TEXT_CHARS, True)


def test_office_files(client: TestClient) -> None:
    as_admin(client)
    document = Document()
    document.add_paragraph("Quarterly report")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Revenue", "42"
    buffer = io.BytesIO()
    document.save(buffer)
    assert attach(client, "report.docx", buffer.getvalue()).json()["text"] == "Quarterly report\nRevenue | 42"

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Activity", None, "Action"])
    sheet.append(["deploy", "api", None])
    buffer = io.BytesIO()
    workbook.save(buffer)
    assert attach(client, "plan.xlsx", buffer.getvalue()).json()["text"] == "Activity |  | Action\ndeploy | api"


def test_unreadable_files_are_refused_with_a_reason(client: TestClient) -> None:
    as_admin(client)
    cases = {
        ("movie.mp4", b"\x00\x01"): "Can't read '.mp4' files",
        ("empty.txt", b""): None,  # refused by the model (min_length) before reading
        ("blank.txt", b"   \n "): "No readable text found in this file.",
        ("binary.txt", bytes(range(128, 256)) * 10): "This doesn't look like a text file",
        ("broken.pdf", b"%PDF-1.4 garbage"): "Could not read this PDF",
    }
    for (filename, content), message in cases.items():
        response = attach(client, filename, content)
        if message is None:
            assert response.status_code == 422
            continue
        assert response.status_code == 400, (filename, response.text)
        assert response.json()["detail"].startswith(message), response.json()

    bad = client.post("/api/attachments/text", json={"filename": "a.txt", "data": "not base64!"})
    assert (bad.status_code, bad.json()["detail"]) == (400, "data must be base64")
